# pylint: disable=redefined-outer-name
# pylint: disable=unused-argument
# pylint: disable=unused-variable
# pylint: disable=too-many-arguments


import datetime
import random
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pytest
from faker import Faker
from openpyxl import Workbook
from simcore_service_webserver.exporter._formatter.xlsx.code_description import (
    CodeDescriptionModel,
    CodeDescriptionParams,
    CodeDescriptionXLSXDocument,
    InputsEntryModel,
    OutputsEntryModel,
    RRIDEntry,
)
from simcore_service_webserver.exporter._formatter.xlsx.dataset_description import (
    ContributorEntryModel,
    DatasetDescriptionParams,
    DatasetDescriptionXLSXDocument,
    DoiEntryModel,
    LinkEntryModel,
)
from simcore_service_webserver.exporter._formatter.xlsx.directory_manifest import (
    DirectoryManifestParams,
    DirectoryManifestXLSXDocument,
    FileEntryModel,
)
from simcore_service_webserver.exporter._formatter.xlsx.submission import (
    SubmissionDocumentParams,
    SubmissionXLSXDocument,
)

MAX_ENTRIES_IN_ARRAYS = 10


# fixtures below


@pytest.fixture
def temp_dir(tmpdir) -> Path:
    yield Path(tmpdir)


# utils below


def get_workbook(xls_path: Path) -> Workbook:
    return openpyxl.load_workbook(xls_path)


def assert_expected_layout(
    workbook: Workbook, expected_layout: dict[str, dict[str, Any]]
) -> bool:
    for sheet_name in expected_layout:
        sheet = workbook[sheet_name]
        sheet_expected_values = expected_layout[sheet_name]
        for address, value in sheet_expected_values.items():
            sheet_value = sheet[address].value
            assert (
                sheet_value == value
            ), f"'{address} value' \"{sheet_value}\" != 'expected value' \"{value}\""

    return True


def random_text(prefix: str = "") -> str:
    prefix_str = f"{prefix}__" if prefix else ""
    return prefix_str + Faker().text()


def column_generator(start_letter: str, elements: int) -> Iterable[str]:
    """
    will only work with a low amount of columns that's why
    MAX_ENTRIES_IN_ARRAYS is a low number
    """
    for k in range(elements):
        yield chr(ord(start_letter) + k)


# tests below


@pytest.mark.parametrize(
    "inputs",
    [
        {
            "award_number": "",
            "milestone_archived": "",
            "milestone_completion_date": None,
        }
    ],
)
def test_submission_document_params(inputs: dict[str, Any]):
    params = SubmissionDocumentParams(**inputs)
    assert params.award_number == ""
    assert params.milestone_archived == ""
    assert params.milestone_completion_date == ""


def test_code_submission(temp_dir: Path):
    award_number = random_text()
    milestone_archived = random_text()
    milestone_completion_date = datetime.datetime.now()

    submission_params = SubmissionDocumentParams(
        award_number=award_number,
        milestone_archived=milestone_archived,
        milestone_completion_date=milestone_completion_date,
    )
    submission_xlsx = SubmissionXLSXDocument()
    submission_xlsx.save_document(base_path=temp_dir, template_data=submission_params)
    workbook = get_workbook(submission_xlsx.document_path(base_path=temp_dir))

    expected_layout = {
        "Sheet1": {
            "C2": award_number,
            "C3": milestone_archived,
            "C4": milestone_completion_date.strftime("%d/%m/%Y"),
        }
    }

    assert assert_expected_layout(workbook, expected_layout) is True


def test_dataset_description(temp_dir: Path):
    name = random_text()
    description = random_text()
    keywords = random_text()
    funding = random_text()

    number_of_subjects: int = random.randint(1, 100)
    number_of_samples: int = random.randint(101, 200)
    completeness_of_dataset = random_text()
    parent_dataset_id = random_text()
    title_for_complete_dataset = random_text()

    contributor_entries = [
        ContributorEntryModel(
            contributor=random_text(f"contributor{i}"),
            orcid_id=random_text(f"orcid_id{i}"),
            description=random_text(f"description{i}"),
            affiliation=random_text(f"affiliation{i}"),
            role=random_text(f"role{i}"),
            is_contact_person=random_text(f"is_contact_person{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]

    doi_entries = [
        DoiEntryModel(
            originating_article_doi=random_text(f"originating_article_doi{i}"),
            protocol_url_or_doi=random_text(f"protocol_url_or_doi{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]

    link_entries = [
        LinkEntryModel(
            additional_link=random_text(f"additional_link{i}"),
            link_description=random_text(f"link_description{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]

    dataset_description_params = DatasetDescriptionParams(
        name=name,
        description=description,
        keywords=keywords,
        contributor_entries=contributor_entries,
        funding=funding,
        doi_entries=doi_entries,
        link_entries=link_entries,
        number_of_subjects=number_of_subjects,
        number_of_samples=number_of_samples,
        completeness_of_dataset=completeness_of_dataset,
        parent_dataset_id=parent_dataset_id,
        title_for_complete_dataset=title_for_complete_dataset,
    )

    dataset_description_xlsx = DatasetDescriptionXLSXDocument()
    dataset_description_xlsx.save_document(
        base_path=temp_dir, template_data=dataset_description_params
    )
    workbook = get_workbook(dataset_description_xlsx.document_path(base_path=temp_dir))

    expected_layout = {
        "Sheet1": {
            "D2": name,
            "D3": description,
            "D4": keywords,
            "D10": "Thank you everyone!",
            "D11": funding,
            "D16": number_of_subjects,
            "D17": number_of_samples,
            "D18": completeness_of_dataset,
            "D19": parent_dataset_id,
            "D20": title_for_complete_dataset,
            "D21": "1.2.3",
        }
    }

    expected_sheet1 = expected_layout["Sheet1"]

    contributor_entry: ContributorEntryModel
    for column_letter, contributor_entry in zip(
        column_generator("D", len(contributor_entries)), contributor_entries
    ):
        expected_sheet1[f"{column_letter}5"] = contributor_entry.contributor
        expected_sheet1[f"{column_letter}6"] = contributor_entry.orcid_id
        expected_sheet1[f"{column_letter}7"] = contributor_entry.affiliation
        expected_sheet1[f"{column_letter}8"] = contributor_entry.role
        expected_sheet1[f"{column_letter}9"] = contributor_entry.is_contact_person

    doi_entry: DoiEntryModel
    for column_letter, doi_entry in zip(
        column_generator("D", len(doi_entries)), doi_entries
    ):
        expected_sheet1[f"{column_letter}12"] = doi_entry.originating_article_doi
        expected_sheet1[f"{column_letter}13"] = doi_entry.protocol_url_or_doi

    link_entry: LinkEntryModel
    for column_letter, link_entry in zip(
        column_generator("D", len(link_entries)), link_entries
    ):
        expected_sheet1[f"{column_letter}14"] = link_entry.additional_link
        expected_sheet1[f"{column_letter}15"] = link_entry.link_description

    assert assert_expected_layout(workbook, expected_layout) is True


def test_code_description(temp_dir: Path):
    # pylint: disable=too-many-statements
    rrid_entires = [
        RRIDEntry(
            rrid_term=random_text(f"rrid_term{i}"),
            rrid_identifier=random_text(f"rrod_identifier{i}"),
            ontological_term=random_text(f"ontological_term{i}"),
            ontological_identifier=random_text(f"ontological_identifier{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]

    tsr1_rating = random.randint(1, 10)
    tsr1_reference = random_text()
    tsr2_rating = random.randint(11, 20)
    tsr2_reference = random_text()
    tsr3_rating = random.randint(21, 30)
    tsr3_reference = random_text()
    tsr4_rating = random.randint(31, 40)
    tsr4_reference = random_text()
    tsr5_rating = random.randint(41, 50)
    tsr5_reference = random_text()
    tsr6_rating = random.randint(51, 60)
    tsr6_reference = random_text()
    tsr7_rating = random.randint(61, 70)
    tsr7_reference = random_text()
    tsr8_rating = random.randint(71, 80)
    tsr8_reference = random_text()
    tsr9_rating = random.randint(81, 90)
    tsr9_reference = random_text()
    tsr10a_rating = random.randint(91, 100)
    tsr10a_reference = random_text()
    tsr10b_relevant_standards = random_text()

    ann1_status = random_text()
    ann1_reference = random_text()
    ann2_status = random_text()
    ann2_reference = random_text()
    ann3_status = random_text()
    ann3_reference = random_text()
    ann4_status = random_text()
    ann4_reference = random_text()
    ann5_status = random_text()
    ann5_reference = random_text()

    reppresentation_in_cell_ml = random_text()

    code_description = CodeDescriptionModel(
        rrid_entires=rrid_entires,
        tsr1_rating=tsr1_rating,
        tsr1_reference=tsr1_reference,
        tsr2_rating=tsr2_rating,
        tsr2_reference=tsr2_reference,
        tsr3_rating=tsr3_rating,
        tsr3_reference=tsr3_reference,
        tsr4_rating=tsr4_rating,
        tsr4_reference=tsr4_reference,
        tsr5_rating=tsr5_rating,
        tsr5_reference=tsr5_reference,
        tsr6_rating=tsr6_rating,
        tsr6_reference=tsr6_reference,
        tsr7_rating=tsr7_rating,
        tsr7_reference=tsr7_reference,
        tsr8_rating=tsr8_rating,
        tsr8_reference=tsr8_reference,
        tsr9_rating=tsr9_rating,
        tsr9_reference=tsr9_reference,
        tsr10a_rating=tsr10a_rating,
        tsr10a_reference=tsr10a_reference,
        tsr10b_relevant_standards=tsr10b_relevant_standards,
        ann1_status=ann1_status,
        ann1_reference=ann1_reference,
        ann2_status=ann2_status,
        ann2_reference=ann2_reference,
        ann3_status=ann3_status,
        ann3_reference=ann3_reference,
        ann4_status=ann4_status,
        ann4_reference=ann4_reference,
        ann5_status=ann5_status,
        ann5_reference=ann5_reference,
        reppresentation_in_cell_ml=reppresentation_in_cell_ml,
    )

    inputs = [
        InputsEntryModel(
            service_alias=random_text(f"service_alias{i}"),
            service_name=random_text(f"service_name{i}"),
            service_version="1.2.3",
            input_name=random_text(f"input_name{i}"),
            input_parameter_description=random_text(f"input_name{i}"),
            input_data_type=random_text(f"input_data_type{i}"),
            input_data_units=random_text(f"input_data_units{i}"),
            input_data_default_value=random_text(f"input_data_default_value{i}"),
            input_data_constraints=random_text(f"input_name{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]
    outputs = [
        OutputsEntryModel(
            service_alias=random_text(f"service_alias{i}"),
            service_name=random_text(f"service_name{i}"),
            service_version="1.2.3",
            output_name=random_text(f"output_name{i}"),
            output_data_ontology_identifier=random_text(
                f"output_data_ontology_identifier{i}"
            ),
            output_data_type=random_text(f"output_data_type{i}"),
            output_data_units=random_text(f"output_data_units{i}"),
        )
        for i in range(MAX_ENTRIES_IN_ARRAYS)
    ]

    dataset_description_params = CodeDescriptionParams(
        code_description=code_description,
        inputs=inputs,
        outputs=outputs,
    )

    code_description_xlsx = CodeDescriptionXLSXDocument()
    code_description_xlsx.save_document(
        base_path=temp_dir, template_data=dataset_description_params
    )
    workbook = get_workbook(code_description_xlsx.document_path(base_path=temp_dir))

    expected_layout = {
        "Code Description": {
            "D7": tsr1_rating,
            "D8": tsr1_reference,
            "D9": tsr2_rating,
            "D10": tsr2_reference,
            "D11": tsr3_rating,
            "D12": tsr3_reference,
            "D13": tsr4_rating,
            "D14": tsr4_reference,
            "D15": tsr5_rating,
            "D16": tsr5_reference,
            "D17": tsr6_rating,
            "D18": tsr6_reference,
            "D19": tsr7_rating,
            "D20": tsr7_reference,
            "D21": tsr8_rating,
            "D22": tsr8_reference,
            "D23": tsr9_rating,
            "D24": tsr9_reference,
            "D25": tsr10a_rating,
            "D26": tsr10a_reference,
            "D27": tsr10b_relevant_standards,
            "D29": ann1_status,
            "D30": ann1_reference,
            "D31": ann2_status,
            "D32": ann2_reference,
            "D33": ann3_status,
            "D34": ann3_reference,
            "D35": ann4_status,
            "D36": ann4_reference,
            "D37": ann5_status,
            "D38": ann5_reference,
            "D41": reppresentation_in_cell_ml,
        },
        "Inputs": {},
        "Outputs": {},
    }

    expected_code_description = expected_layout["Code Description"]
    for column_letter, rrid_entry in zip(
        column_generator("D", len(rrid_entires)), rrid_entires
    ):
        rrid_entry: RRIDEntry = rrid_entry

        expected_code_description[f"{column_letter}2"] = rrid_entry.rrid_term
        expected_code_description[f"{column_letter}3"] = rrid_entry.rrid_identifier
        expected_code_description[f"{column_letter}4"] = rrid_entry.ontological_term
        expected_code_description[
            f"{column_letter}5"
        ] = rrid_entry.ontological_identifier

    expected_inputs = expected_layout["Inputs"]
    for row, input_entry in zip(range(4, len(inputs) + 4), inputs):
        input_entry: InputsEntryModel = input_entry

        expected_inputs[f"B{row}"] = input_entry.service_alias
        expected_inputs[f"C{row}"] = input_entry.service_name
        expected_inputs[f"D{row}"] = input_entry.service_version
        expected_inputs[f"E{row}"] = input_entry.input_name
        expected_inputs[f"F{row}"] = input_entry.input_parameter_description
        expected_inputs[f"G{row}"] = input_entry.input_data_type
        expected_inputs[f"H{row}"] = input_entry.input_data_units
        expected_inputs[f"I{row}"] = input_entry.input_data_default_value

    expected_outputs = expected_layout["Outputs"]
    for row, output_entry in zip(range(4, len(outputs) + 4), outputs):
        output_entry: OutputsEntryModel = output_entry

        expected_outputs[f"B{row}"] = output_entry.service_alias
        expected_outputs[f"C{row}"] = output_entry.service_name
        expected_outputs[f"D{row}"] = output_entry.service_version
        expected_outputs[f"E{row}"] = output_entry.output_name
        expected_outputs[f"F{row}"] = output_entry.output_data_ontology_identifier
        expected_outputs[f"G{row}"] = output_entry.output_data_type
        expected_outputs[f"H{row}"] = output_entry.output_data_units

    assert assert_expected_layout(workbook, expected_layout) is True


def test_directory_manifest(temp_dir: Path, dir_with_random_content: Path):
    # generate dir structure of which to have file data

    directory_manifest_params = DirectoryManifestParams.compose_from_directory(
        dir_with_random_content
    )

    # inject random descrition & additional_metadata
    file_entry: FileEntryModel
    for k, file_entry in enumerate(directory_manifest_params.file_entries):
        file_entry.description = random_text(f"description{k}")
        file_entry.additional_metadata = [
            random_text(f"{k}_{i}_additional_metadata")
            for i in range(MAX_ENTRIES_IN_ARRAYS)
        ]

    directory_manifest_xlsx = DirectoryManifestXLSXDocument()
    directory_manifest_xlsx.save_document(
        base_path=temp_dir, template_data=directory_manifest_params
    )
    workbook = get_workbook(directory_manifest_xlsx.document_path(base_path=temp_dir))

    expected_layout = {"Sheet1": {}}

    expected_sheet1 = expected_layout["Sheet1"]
    file_entry: FileEntryModel
    for row, file_entry in zip(
        range(2, len(directory_manifest_params.file_entries) + 2),
        directory_manifest_params.file_entries,
    ):
        expected_sheet1[f"C{row}"] = file_entry.description
        expected_sheet1[f"A{row}"] = file_entry.filename
        expected_sheet1[f"B{row}"] = file_entry.timestamp
        expected_sheet1[f"D{row}"] = file_entry.file_type
        # write down additional_metadata
        for k, column_letter, metadata_string in zip(
            range(len(file_entry.additional_metadata)),
            column_generator("E", len(file_entry.additional_metadata)),
            file_entry.additional_metadata,
        ):
            expected_sheet1[f"{column_letter}{row}"] = metadata_string

    assert assert_expected_layout(workbook, expected_layout) is True
