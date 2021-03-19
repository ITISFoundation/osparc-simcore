from simcore_service_webserver.exporter.xlsx_base import (
    BaseXLSXCellData,
    BaseXLSXSheet,
    BaseXLSXDocument,
)

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.styles.borders import BORDER_THIN, BORDER_MEDIUM


COLOR_BLACK = "FF000000"
COLOR_LINK = "0563C1"
COLOR_BLUE = "8CB3DF"
COLOR_GREEN = "99C87B"
COLOR_LIGHT_GREEN = "9AC87C"
COLOR_YELLOW = "FEE288"


class T(BaseXLSXCellData):
    """Helper for inputing text into cells"""

    alignment = Alignment(wrap_text=True)

    def __init__(self, text: str):
        super().__init__(value=text)


class TB(T):
    font = Font(bold=True)


class Link(T):
    font = Font(underline="single", color=COLOR_LINK)

    def __init__(self, text: str, link: str):
        super().__init__(text=f'=HYPERLINK("{link}", "{text}")')


class BackgroundWithColor(BaseXLSXCellData):
    def __init__(self, color: str):
        super().__init__(
            fill=PatternFill(start_color=color, end_color=color, fill_type="solid")
        )


class Backgrounds:
    blue = BackgroundWithColor(color=COLOR_BLUE)
    green = BackgroundWithColor(color=COLOR_GREEN)
    light_green = BackgroundWithColor(color=COLOR_LIGHT_GREEN)
    yellow = BackgroundWithColor(color=COLOR_YELLOW)


class BorderWithStyle(BaseXLSXCellData):
    """Base for creating custom borders"""

    def __init__(self, *borders_sides, border_style: str, color: str):
        side = Side(border_style=border_style, color=color)
        super().__init__(border=Border(**{x: side for x in borders_sides}))


class Borders:
    """Collector for different border styles"""

    light_grid = BorderWithStyle(
        "top",
        "left",
        "right",
        "bottom",
        "outline",
        "vertical",
        "horizontal",
        border_style=BORDER_THIN,
        color=COLOR_BLACK,
    )

    bold_grid = BorderWithStyle(
        "top",
        "left",
        "right",
        "bottom",
        "outline",
        "vertical",
        "horizontal",
        border_style=BORDER_MEDIUM,
        color=COLOR_BLACK,
    )

    border_bottom_thick = BorderWithStyle(
        "bottom", border_style=BORDER_MEDIUM, color=COLOR_BLACK
    )
    border_bottom_thick = BorderWithStyle(
        "bottom", border_style=BORDER_MEDIUM, color=COLOR_BLACK
    )
    border_left_thick = BorderWithStyle(
        "left", border_style=BORDER_MEDIUM, color=COLOR_BLACK
    )
    border_right_thick = BorderWithStyle(
        "right", border_style=BORDER_MEDIUM, color=COLOR_BLACK
    )


class AllignTopCenter(BaseXLSXCellData):
    alignment = Alignment(horizontal="center", vertical="top")


class AllignTop(BaseXLSXCellData):
    alignment = Alignment(vertical="top")


class SheetOne(BaseXLSXSheet):
    name = "Code Description"
    cell_styles = {
        ## Header
        "A1": TB("Metadata element"),
        "B1": TB("Description"),
        "C1": TB("Example"),
        "A1:C1": Backgrounds.blue,
        ## Classifiers section
        "A2": TB("RRID Term"),
        "B2": T("Associated tools or resources used"),
        "C2": T("ImageJ"),
        "A3": TB("RRID Identifier"),
        "B3": T("Associated tools or resources identifier (with 'RRID:')"),
        "C3": T("RRID:SCR_003070"),
        "A4": TB("Ontological term"),
        "B4": T("Associated ontological term (human-readable)"),
        "C4": T("Heart"),
        "A5": TB("Ontological Identifier"),
        "B5": Link(
            "Associated ontological identifier from SciCrunch",
            "https://scicrunch.org/sawg",
        ),
        "C5": T("UBERON:0000948"),
        "A2:C5": Backgrounds.green,
        "A5:C5": Borders.border_bottom_thick,
        "A1:C5": Borders.light_grid,
        ## TSR section
        "A6": Link(
            "Ten Simple Rules (TSR)",
            "https://www.imagwiki.nibib.nih.gov/content/10-simple-rules-conformance-rubric",
        ),
        "B7": T(
            "The TSR is a communication tool for modelers to organize their model development process and present it coherently."
        ),
        "C8": Link("Rating (0-4)", "#'TSR Rating Rubric'!A1"),
    }


class SheetTSRRating(BaseXLSXSheet):
    name = "TSR Rating Rubric"
    cell_styles = {
        "A1": T("Conformance Level"),
        "A3": T("Description"),
        "B1": T("Comprehensive"),
        "B2": T(4),
        "B3": T(
            "Can be understood by non MS&P practitioners familiar with the application domain and the intended context of use"
        ),
        "C1": T("Extensive"),
        "C2": T(3),
        "C3": T(
            "Can be understood by MS&P practitions not familiar with the application domain and the intended context of use"
        ),
        "D1": T("Adequate"),
        "D2": T(2),
        "D3": T(
            "Can be understood by MS&P practitioners familiar with the application domain and the intended context of use"
        ),
        "E1": T("Partial"),
        "E2": T(1),
        "E3": T(
            "Unclear to the MS&P practitioners familiar with the application domain and the intended context of use"
        ),
        "F1": T("Insufficient"),
        "F2": T(0),
        "F3": T(
            "Missing or grossly incomplete information to properly evaluate the conformance with the rule"
        ),
        # background
        "A1:F2": Backgrounds.green | AllignTopCenter(),
        "A3:F3": Backgrounds.yellow,
        # Borders
        "A1:F3": Borders.light_grid,
        # Alignment
        #"A1:F2": AllignTopCenter(),
        "A3:F3": AllignTop(),
    }
    cell_merge = {"A1:A2"}
    #column_dimensions = {"A": 20, "B": 20, "C": 20, "D": 20, "E": 20, "F": 20}


class XLSXTestDocument(BaseXLSXDocument):
    sheet_1 = SheetOne()
    sheet_2 = SheetTSRRating()


if __name__ == "__main__":
    document = XLSXTestDocument()
    # document = BaseXLSXDocument(SheetOne(), SecondSheet())
    print(document)
    document.save_document("test.xlsx")
